"""Read the inspected XTB XLSX format without modifying the workbook."""

import hashlib
import re
from datetime import UTC, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel
from pydantic import ValidationError

from trading_assistant.domain.portfolio import (
    BrokerReport,
    CashEvent,
    ImportValidationError,
    InstrumentData,
    ReportPosition,
)

TRADE_COMMENT = re.compile(
    r"^(OPEN|CLOSE) BUY (\d+(?:\.\d+)?)(?:/(\d+(?:\.\d+)?))? @ (\d+(?:\.\d+)?)$"
)
TOLERANCE = Decimal("0.00000001")
KNOWN_CASH_TYPES = {
    "Stock purchase",
    "Stock sell",
    "Subaccount transfer",
    "Deposit",
    "Withdrawal",
    "Dividend",
    "Withholding tax",
    "Free funds interest",
    "Free funds interest tax",
    "Transfer",
    "Tax IFTT",
}


def text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def number(value) -> Decimal:
    try:
        result = Decimal(str(value))
    except (InvalidOperation, ValueError) as exc:
        raise ImportValidationError("Missing or invalid numeric report field.") from exc
    if not result.is_finite():
        raise ImportValidationError("Non-finite numeric report field.")
    return result


def included(product) -> bool:
    label = text(product).casefold()
    if label == "my trades":
        return True
    if label in {"investment plan", "investment plans"}:
        return False
    raise ImportValidationError("Missing or unknown Product label; scope cannot be determined.")


def headers(rows: list[tuple], required: set[str]) -> tuple[int, dict[str, int]]:
    for idx, row in enumerate(rows):
        mapping = {text(value): col for col, value in enumerate(row) if value is not None}
        if required <= mapping.keys():
            return idx, mapping
    raise ImportValidationError(
        f"Report table is missing required columns: {', '.join(sorted(required))}"
    )


def field(row, mapping, name):
    return row[mapping[name]] if mapping[name] < len(row) else None


def date(value, epoch) -> datetime:
    if isinstance(value, (float, int)):
        value = from_excel(value, epoch)
    if not isinstance(value, datetime):
        raise ImportValidationError("Missing or invalid report timestamp.")
    return value.replace(tzinfo=UTC) if value.tzinfo is None else value.astimezone(UTC)


def metadata(rows, label):
    for row in rows:
        if row and text(row[0]) == label:
            return row[1] if len(row) > 1 else None
    raise ImportValidationError(f"Missing report metadata: {label}")


def raw_record(row, mapping, source_row: int) -> dict:
    return {"source_row": source_row, **{key: text(field(row, mapping, key)) for key in mapping}}


def account_currency(sheets: dict[str, list[tuple]]) -> str:
    """Require explicit account denomination, never infer it from a ticker."""
    currencies = set()
    for rows in sheets.values():
        for row in rows:
            if row and text(row[0]).casefold() in {"account currency", "base currency"}:
                currencies.add(text(row[1] if len(row) > 1 else None).upper())

    # The inspected export identifies account-value currency in this summary.
    rows = sheets["Open Positions"]
    for idx, row in enumerate(rows):
        mapping = {text(value): col for col, value in enumerate(row) if value is not None}
        if {"Product", "Metric", "Amount", "Currency"} <= mapping.keys():
            for summary in rows[idx + 1 :]:
                if "Instrument/Position" in summary:
                    break
                if text(field(summary, mapping, "Product")).casefold() == "my trades":
                    currencies.add(text(field(summary, mapping, "Currency")).upper())
            break
    if not currencies or "" in currencies:
        raise ImportValidationError("Account currency cannot be established from the report.")
    if currencies != {"PLN"}:
        raise ImportValidationError(
            "Report account currency is conflicting or unsupported; only verified PLN is supported."
        )
    return "PLN"


class XtbReportReader:
    def __init__(self, account_number: str | None = None):
        self.account_number = account_number

    def read(self, path: Path) -> BrokerReport:
        try:
            data = path.read_bytes()
            workbook = load_workbook(BytesIO(data), read_only=True, data_only=False)
        except Exception as exc:
            raise ImportValidationError("Unable to read XLSX broker report.") from exc
        try:
            return self._read(workbook, hashlib.sha256(data).hexdigest())
        except (ValidationError, KeyError, IndexError) as exc:
            raise ImportValidationError("Invalid broker report structure or values.") from exc
        finally:
            workbook.close()

    def _read(self, workbook, fingerprint: str) -> BrokerReport:
        names = ("Cash Operations", "Closed Positions", "Open Positions")
        if not set(names) <= set(workbook.sheetnames):
            raise ImportValidationError(
                "Expected Cash Operations, Closed Positions, and Open Positions."
            )
        sheets = {name: list(workbook[name].values) for name in names}
        accounts = {text(metadata(rows, "Account number")) for rows in sheets.values()} - {""}
        if len(accounts) > 1 or (
            accounts and self.account_number and accounts != {self.account_number}
        ):
            raise ImportValidationError(
                "Report account metadata conflicts with the selected account."
            )
        account = self.account_number or next(iter(accounts), "")
        if not account:
            raise ImportValidationError("Account metadata is blank. Supply --account explicitly.")
        currency = account_currency(sheets)
        as_of = date(
            metadata(sheets["Open Positions"], "Data as of report generated"), workbook.epoch
        )
        instruments: dict[str, InstrumentData] = {}
        warnings = ["Account currency PLN; instrument quote currencies have not yet been resolved."]
        excluded = 0

        def instrument(symbol, name, category):
            symbol, name, category = text(symbol), text(name), text(category)
            if not symbol:
                raise ImportValidationError("An instrument row has no broker ticker.")
            asset = category if category in {"STOCK", "ETF", "ETC"} else "UNKNOWN"
            item = InstrumentData(broker_symbol=symbol, name=name or symbol, asset_type=asset)
            previous = instruments.get(symbol)
            if previous and previous.asset_type != item.asset_type:
                raise ImportValidationError("Conflicting instrument categories within report.")
            instruments[symbol] = item
            if asset == "UNKNOWN":
                warnings.append(
                    f"Unverified instrument category for {symbol}; AI analysis unavailable."
                )

        rows = sheets["Cash Operations"]
        idx, mapping = headers(
            rows,
            {
                "Type",
                "Instrument",
                "Ticker",
                "Category",
                "Time",
                "Amount",
                "ID",
                "Comment",
                "Product",
                "Position ID",
            },
        )
        events = []
        seen = set()
        for source_row, row in enumerate(rows[idx + 1 :], idx + 2):
            kind = text(field(row, mapping, "Type"))
            if not kind or kind.casefold() == "total":
                continue
            if not included(field(row, mapping, "Product")):
                excluded += 1
                continue
            if "Currency" in mapping and text(field(row, mapping, "Currency")).upper() != currency:
                raise ImportValidationError("Cash operation currency differs from verified PLN.")
            operation_id = text(field(row, mapping, "ID"))
            if not operation_id or operation_id in seen:
                raise ImportValidationError("Missing or repeated cash operation ID within report.")
            seen.add(operation_id)
            symbol = text(field(row, mapping, "Ticker")) or None
            if symbol:
                instrument(
                    symbol, field(row, mapping, "Instrument"), field(row, mapping, "Category")
                )
            event = CashEvent(
                operation_id=operation_id,
                kind=kind,
                occurred_at=date(field(row, mapping, "Time"), workbook.epoch),
                amount=number(field(row, mapping, "Amount")),
                symbol=symbol,
                position_id=text(field(row, mapping, "Position ID")) or None,
                comment=text(field(row, mapping, "Comment")),
                source_row=source_row,
            )
            if event.occurred_at > as_of:
                raise ImportValidationError("Cash event occurs after the report snapshot.")
            if kind in {"Stock purchase", "Stock sell"}:
                match = TRADE_COMMENT.fullmatch(event.comment)
                if not match or not event.symbol or not event.position_id:
                    raise ImportValidationError(
                        f"Unsupported trade comment or identifiers at cash row {source_row}."
                    )
                opening = kind == "Stock purchase"
                if match[1] != ("OPEN" if opening else "CLOSE"):
                    raise ImportValidationError("Trade type and comment direction disagree.")
                if (opening and event.amount >= 0) or (not opening and event.amount <= 0):
                    raise ImportValidationError("Trade cash amount has an unexpected sign.")
                event = event.model_copy(
                    update={
                        "action": "BUY" if opening else "SELL",
                        "quantity": number(match[2]),
                        "execution_price": number(match[4]),
                    }
                )
                event = CashEvent.model_validate(event.model_dump())
                if match[3] and event.quantity > number(match[3]):
                    raise ImportValidationError("Execution quantity exceeds contextual total.")
            if kind not in KNOWN_CASH_TYPES:
                warnings.append(f"Unrecognized cash event type retained for review: {kind}")
            events.append(event)

        rows = sheets["Closed Positions"]
        idx, mapping = headers(
            rows,
            {
                "Instrument",
                "Ticker",
                "Category",
                "Product",
                "Position ID",
                "Volume",
                "Type",
            },
        )
        closed = []
        for source_row, row in enumerate(rows[idx + 1 :], idx + 2):
            if not text(field(row, mapping, "Ticker")):
                continue
            if not included(field(row, mapping, "Product")):
                excluded += 1
                continue
            if text(field(row, mapping, "Type")) != "BUY":
                raise ImportValidationError(
                    "Only non-leveraged long holdings are supported initially."
                )
            instrument(
                field(row, mapping, "Ticker"),
                field(row, mapping, "Instrument"),
                field(row, mapping, "Category"),
            )
            closed.append(raw_record(row, mapping, source_row))

        rows = sheets["Open Positions"]
        idx, mapping = headers(
            rows,
            {
                "Product",
                "Instrument/Position",
                "Ticker",
                "Category",
                "Type",
                "Volume",
                "Value",
                "Net Profit %",
            },
        )
        positions = []
        current = None
        lot_quantity = Decimal("0")

        def finish():
            if current is not None and abs(current.quantity - lot_quantity) > TOLERANCE:
                raise ImportValidationError(
                    "Open instrument total disagrees with its detailed lots."
                )

        for source_row, row in enumerate(rows[idx + 1 :], idx + 2):
            symbol = text(field(row, mapping, "Ticker"))
            if not symbol:
                continue
            if not included(field(row, mapping, "Product")):
                excluded += 1
                continue
            category = text(field(row, mapping, "Category"))
            if category:
                finish()
                instrument(symbol, field(row, mapping, "Instrument/Position"), category)
                pnl = field(row, mapping, "Net Profit %")
                current = ReportPosition(
                    symbol=symbol,
                    quantity=number(field(row, mapping, "Volume")),
                    market_value=number(field(row, mapping, "Value")),
                    pnl_percent=number(pnl) if pnl is not None else None,
                    source_row=source_row,
                )
                positions.append(current)
                lot_quantity = Decimal("0")
            else:
                if (
                    current is None
                    or current.symbol != symbol
                    or text(field(row, mapping, "Type")) != "BUY"
                ):
                    raise ImportValidationError("Unrecognized open-position lot hierarchy.")
                quantity = number(field(row, mapping, "Volume"))
                if quantity <= 0:
                    raise ImportValidationError("Open lot quantity must be positive.")
                lot_quantity += quantity
                current.lots.append(raw_record(row, mapping, source_row))
        finish()
        if len({p.symbol for p in positions}) != len(positions):
            raise ImportValidationError(
                "Repeated My Trades instrument summaries require explicit grouping."
            )
        return BrokerReport(
            account_number=account,
            currency=currency,
            fingerprint=fingerprint,
            as_of=as_of,
            instruments=list(instruments.values()),
            cash_events=events,
            positions=positions,
            closed_records=closed,
            excluded_rows=excluded,
            warnings=sorted(set(warnings)),
        )
