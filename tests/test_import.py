from datetime import UTC, datetime
from decimal import Decimal

import pytest
from openpyxl import Workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from trading_assistant.adapters.database import (
    Account,
    CashOperation,
    ImportBatch,
    Instrument,
    Transaction,
    database_engine,
    migrate,
)
from trading_assistant.adapters.portfolio_repository import SqlAlchemyPortfolioRepository
from trading_assistant.adapters.xtb import XtbReportReader
from trading_assistant.application.portfolio import PortfolioDashboardService
from trading_assistant.domain.portfolio import ImportValidationError


@pytest.fixture
def report_path(tmp_path):
    workbook = Workbook()
    workbook.remove(workbook.active)
    cash = workbook.create_sheet("Cash Operations")
    cash.append(["Account number", "synthetic-account"])
    cash.append(["Cash Operations"])
    cash.append(["Date from (UTC)", datetime(2026, 1, 1)])
    cash.append(["Date to (UTC)", datetime(2026, 10, 1)])
    cash.append(
        [
            "Type",
            "Instrument",
            "Ticker",
            "Category",
            "Time",
            "Amount",
            "ID",
            "Comment",
            "Product",
            "Position ID",
        ]
    )
    time = datetime(2026, 9, 1, 10)
    cash.append(["Deposit", None, None, None, time, 1000, "d1", "deposit", "My Trades", None])
    cash.append(
        [
            "Stock purchase",
            "Test Stock",
            "TEST.US",
            "STOCK",
            time,
            -100,
            "b1",
            "OPEN BUY 1/1.5 @ 100",
            "My Trades",
            "101",
        ]
    )
    cash.append(
        [
            "Stock purchase",
            "Test Stock",
            "TEST.US",
            "STOCK",
            time,
            -50,
            "b2",
            "OPEN BUY 0.5/1.5 @ 100",
            "My Trades",
            "101",
        ]
    )
    cash.append(
        [
            "Stock sell",
            "Test Stock",
            "TEST.US",
            "STOCK",
            time,
            25,
            "s1",
            "CLOSE BUY 0.25/1.5 @ 100",
            "My Trades",
            "101",
        ]
    )
    cash.append(
        [
            "Stock purchase",
            "Test ETF",
            "ETF.DE",
            "ETF",
            time,
            -40,
            "b3",
            "OPEN BUY 0.4 @ 100",
            "My Trades",
            "102",
        ]
    )
    cash.append(
        [
            "Subaccount transfer",
            None,
            None,
            None,
            time,
            -50,
            "t1",
            "Transfer to plan",
            "My Trades",
            None,
        ]
    )
    cash.append(
        [
            "Stock purchase",
            "Excluded ETF",
            "EXCLUDED.DE",
            "ETF",
            time,
            -10,
            "x1",
            "unsupported excluded comment",
            "Investment Plans",
            "999",
        ]
    )
    cash.append(["Total", None, None, None, None, 785])
    closed = workbook.create_sheet("Closed Positions")
    closed.append(["Account number", "synthetic-account"])
    closed.append(["Instrument", "Ticker", "Category", "Product", "Position ID", "Volume", "Type"])
    closed.append(["Test Stock", "TEST.US", "STOCK", "My Trades", "101", 0.25, "BUY"])
    closed.append(["Excluded ETF", "EXCLUDED.DE", "ETF", "Investment Plans", "999", 1, "BUY"])
    opened = workbook.create_sheet("Open Positions")
    opened.append(["Account number", "synthetic-account"])
    opened.append(["Data as of report generated", datetime(2026, 10, 1, 12)])
    opened.append(["Product", "Metric", "Amount", "Currency"])
    opened.append(["My Trades", "Open position value", 165, "PLN"])
    opened.append(["Investment Plans", "Open position value", 10, "EUR"])
    opened.append(
        [
            "Product",
            "Instrument/Position",
            "Ticker",
            "Category",
            "Type",
            "Volume",
            "Value",
            "Net Profit %",
        ]
    )
    opened.append(["My Trades", "Test Stock", "TEST.US", "STOCK", None, 1.25, 125, 0])
    opened.append(["My Trades", "101", "TEST.US", None, "BUY", 1.25, 125, 0])
    opened.append(["My Trades", "Test ETF", "ETF.DE", "ETF", None, 0.4, 40, 0])
    opened.append(["My Trades", "102", "ETF.DE", None, "BUY", 0.4, 40, 0])
    opened.append(["Investment Plan", "Excluded ETF", "EXCLUDED.DE", "ETF", None, 1, 10, 0])
    opened.append(["Investment Plan", "999", "EXCLUDED.DE", None, "BUY", 1, 10, 0])
    path = tmp_path / "synthetic.xlsx"
    workbook.save(path)
    return path


@pytest.fixture
def repository(tmp_path):
    engine = database_engine(tmp_path / "portfolio.sqlite3")
    migrate(engine)
    yield SqlAlchemyPortfolioRepository(engine)
    engine.dispose()


def count(repository, model):
    with Session(repository.engine) as session:
        return session.scalar(select(func.count()).select_from(model))


def test_reader_filters_products_and_preserves_fractional_events(report_path):
    report = XtbReportReader().read(report_path)
    assert len(report.cash_events) == 6
    assert report.excluded_rows == 4
    assert len(report.positions) == 2
    assert report.cash_events[2].quantity == Decimal("0.5")
    assert report.cash_events[3].quantity == Decimal("0.25")
    assert {i.asset_type for i in report.instruments} == {"STOCK", "ETF"}
    assert all(i.broker_symbol != "EXCLUDED.DE" for i in report.instruments)


def test_import_derives_holdings_and_does_not_double_count_cash(repository, report_path):
    report = XtbReportReader().read(report_path)
    result = repository.import_report(report)
    assert (result.cash_added, result.trades_added) == (6, 4)
    assert count(repository, Transaction) == 4
    assert count(repository, CashOperation) == 6
    state = repository.portfolio()
    assert {h.symbol: h.quantity for h in state.holdings} == {
        "TEST.US": Decimal("1.25"),
        "ETF.DE": Decimal("0.4"),
    }
    assert state.cash_movement == Decimal("785")
    assert count(repository, Instrument) == 2


def test_repeat_and_overlapping_reports_skip_events(repository, report_path):
    report = XtbReportReader().read(report_path)
    repository.import_report(report)
    assert repository.import_report(report).already_imported
    overlapping = report.model_copy(update={"fingerprint": "a" * 64})
    result = repository.import_report(overlapping)
    assert result.cash_duplicates == 6
    assert result.cash_added == 0
    assert count(repository, Transaction) == 4
    assert count(repository, ImportBatch) == 2


def test_conflicting_event_rolls_back_entire_batch(repository, report_path):
    report = XtbReportReader().read(report_path)
    repository.import_report(report)
    conflicting = report.model_copy(deep=True)
    conflicting.fingerprint = "b" * 64
    conflicting.cash_events[0].amount = Decimal("999")
    new_event = report.cash_events[0].model_copy(update={"operation_id": "new-deposit"})
    conflicting.cash_events.insert(0, new_event)
    with pytest.raises(ImportValidationError, match="conflicts"):
        repository.import_report(conflicting)
    assert count(repository, ImportBatch) == 1
    assert count(repository, CashOperation) == 6
    assert repository.portfolio().cash_movement == Decimal("785")


def test_incomplete_history_rolls_back_account_and_all_records(repository, report_path):
    report = XtbReportReader().read(report_path)
    report.cash_events = [e for e in report.cash_events if e.operation_id != "b2"]
    with pytest.raises(ImportValidationError, match="reconcile"):
        repository.import_report(report)
    for model in (Account, Instrument, ImportBatch, CashOperation, Transaction):
        assert count(repository, model) == 0


def test_multi_account_identity_and_decimal_precision(repository, report_path):
    report = XtbReportReader().read(report_path)
    report.cash_events[0].amount = Decimal("1000.123456789012345678")
    repository.import_report(report)
    repository.import_report(report.model_copy(update={"account_number": "another-account"}))
    assert count(repository, Account) == 2
    assert count(repository, CashOperation) == 12
    assert repository.portfolio().cash_movement == Decimal("1570.246913578024691356")
    assert {h.symbol: h.quantity for h in repository.portfolio().holdings}["TEST.US"] == Decimal(
        "2.50"
    )


def test_sqlite_pragmas_and_foreign_keys(repository):
    from sqlalchemy.exc import IntegrityError

    with repository.engine.connect() as connection:
        assert connection.exec_driver_sql("PRAGMA foreign_keys").scalar() == 1
        assert connection.exec_driver_sql("PRAGMA journal_mode").scalar() == "wal"
    with Session(repository.engine) as session, pytest.raises(IntegrityError):
        session.add(
            ImportBatch(
                account_id=999,
                fingerprint="c" * 64,
                as_of="x",
                excluded_rows=0,
                warnings=[],
                closed_records=[],
            )
        )
        session.commit()


async def test_database_dashboard_uses_report_values_and_keeps_cash_unknown(
    repository, report_path
):
    repository.import_report(XtbReportReader().read(report_path))
    snapshot = await PortfolioDashboardService(repository).snapshot()
    assert len(snapshot.positions) == 2
    assert snapshot.cash is None
    assert snapshot.portfolio_value is None
    assert all(p.market_value_usd is None for p in snapshot.positions)
    assert any("not live quotes" in a for a in snapshot.alerts)
    assert any("unverified" in a for a in snapshot.alerts)


def test_account_override_must_match_report_metadata(report_path):
    with pytest.raises(ImportValidationError, match="conflicts"):
        XtbReportReader("wrong-account").read(report_path)


def test_blank_account_requires_explicit_selection(report_path):
    from openpyxl import load_workbook

    workbook = load_workbook(report_path)
    for sheet in workbook:
        sheet["B1"] = None
    workbook.save(report_path)
    with pytest.raises(ImportValidationError, match="--account"):
        XtbReportReader().read(report_path)
    assert (
        XtbReportReader("selected-account").read(report_path).account_number == "selected-account"
    )


@pytest.mark.parametrize("product", [None, "Unrecognized product"])
def test_unknown_scope_is_not_silently_imported(report_path, product):
    from openpyxl import load_workbook

    workbook = load_workbook(report_path)
    workbook["Cash Operations"]["I6"] = product
    workbook.save(report_path)
    with pytest.raises(ImportValidationError, match="Product"):
        XtbReportReader().read(report_path)


def test_bad_trade_comment_and_duplicate_ids_are_rejected(report_path):
    from openpyxl import load_workbook

    workbook = load_workbook(report_path)
    workbook["Cash Operations"]["H7"] = "unsupported trade format"
    workbook.save(report_path)
    with pytest.raises(ImportValidationError, match="trade comment"):
        XtbReportReader().read(report_path)
    workbook["Cash Operations"]["H7"] = "OPEN BUY 1/1.5 @ 100"
    workbook["Cash Operations"]["G8"] = "b1"
    workbook.save(report_path)
    with pytest.raises(ImportValidationError, match="repeated"):
        XtbReportReader().read(report_path)


def test_migrations_are_repeatable_and_packaged(repository):
    from alembic.autogenerate import compare_metadata
    from alembic.migration import MigrationContext

    from trading_assistant.adapters.database import Base

    migrate(repository.engine)
    with repository.engine.connect() as connection:
        changes = compare_metadata(MigrationContext.configure(connection), Base.metadata)
    assert changes == []


def test_cli_preview_does_not_create_database(report_path, tmp_path):
    from typer.testing import CliRunner

    from trading_assistant.cli import app

    path = tmp_path / "not-created.sqlite3"
    result = CliRunner().invoke(
        app, ["--database", str(path), "import-xtb", str(report_path), "--dry-run"]
    )
    assert result.exit_code == 0
    assert "Database unchanged" in result.output
    assert not path.exists()


@pytest.mark.parametrize("currency", ["EUR", "USD", None])
def test_unverified_currency_is_rejected_before_persistence(repository, report_path, currency):
    from openpyxl import load_workbook

    from trading_assistant.application.portfolio import ImportService

    workbook = load_workbook(report_path)
    workbook["Open Positions"]["D4"] = currency
    workbook.save(report_path)
    with pytest.raises(ImportValidationError, match="currency"):
        ImportService(XtbReportReader(), repository).import_file(report_path)
    assert count(repository, Account) == 0
    assert count(repository, CashOperation) == 0


def test_absent_currency_evidence_is_rejected(report_path):
    from openpyxl import load_workbook

    workbook = load_workbook(report_path)
    workbook["Open Positions"].delete_rows(3, 3)
    workbook.save(report_path)
    with pytest.raises(ImportValidationError, match="currency cannot be established"):
        XtbReportReader().read(report_path)


def test_currency_metadata_must_agree_with_summary(report_path):
    from openpyxl import load_workbook

    workbook = load_workbook(report_path)
    sheet = workbook["Cash Operations"]
    sheet.insert_rows(2)
    sheet["A2"] = "Account currency"
    sheet["B2"] = "USD"
    workbook.save(report_path)
    with pytest.raises(ImportValidationError, match="currency"):
        XtbReportReader().read(report_path)
    sheet["B2"] = "PLN"
    workbook["Open Positions"].delete_rows(3, 3)
    workbook.save(report_path)
    assert XtbReportReader().read(report_path).currency == "PLN"


def test_cash_currency_must_agree_with_account_denomination(report_path):
    from openpyxl import load_workbook

    workbook = load_workbook(report_path)
    sheet = workbook["Cash Operations"]
    sheet["K5"] = "Currency"
    sheet["K6"] = "USD"
    workbook.save(report_path)
    with pytest.raises(ImportValidationError, match="Cash operation currency"):
        XtbReportReader().read(report_path)


def test_historical_sale_cannot_invalidate_latest_valuation(repository, report_path):
    report = XtbReportReader().read(report_path)
    repository.import_report(report)
    historical = report.model_copy(deep=True)
    historical.fingerprint = "d" * 64
    historical.as_of = datetime(2026, 9, 15, tzinfo=UTC)
    sale = report.cash_events[3].model_copy(
        update={
            "operation_id": "historical-sale",
            "occurred_at": datetime(2026, 9, 2, tzinfo=UTC),
        }
    )
    historical.cash_events.append(sale)
    stock = next(p for p in historical.positions if p.symbol == "TEST.US")
    stock.quantity = Decimal("1")
    stock.market_value = Decimal("100")
    stock.lots[0]["Volume"] = "1"
    with pytest.raises(ImportValidationError, match="invalidate a stored portfolio snapshot"):
        repository.import_report(historical)
    assert count(repository, ImportBatch) == 1
    assert count(repository, CashOperation) == 6
    assert count(repository, Transaction) == 4
    latest = next(h for h in repository.portfolio().holdings if h.symbol == "TEST.US")
    assert latest.quantity == Decimal("1.25")
    assert latest.report_value == Decimal("125")


def test_balanced_historical_additions_preserve_existing_snapshots(repository, report_path):
    report = XtbReportReader().read(report_path)
    repository.import_report(report)
    historical = report.model_copy(deep=True)
    historical.fingerprint = "e" * 64
    historical.as_of = datetime(2026, 9, 15, tzinfo=UTC)
    buy = report.cash_events[1].model_copy(
        update={
            "operation_id": "historical-buy",
            "quantity": Decimal("0.25"),
            "amount": Decimal("-25"),
            "comment": "OPEN BUY 0.25 @ 100",
            "occurred_at": datetime(2026, 9, 2, tzinfo=UTC),
        }
    )
    sale = report.cash_events[3].model_copy(
        update={
            "operation_id": "historical-sale",
            "occurred_at": datetime(2026, 9, 3, tzinfo=UTC),
        }
    )
    historical.cash_events.extend([buy, sale])
    result = repository.import_report(historical)
    assert result.trades_added == 2
    latest = next(h for h in repository.portfolio().holdings if h.symbol == "TEST.US")
    assert latest.quantity == Decimal("1.25")
    assert latest.report_value == Decimal("125")
    assert max(repository.portfolio().report_dates) == report.as_of


def test_new_report_checks_affected_snapshots_before_its_own_date(repository, report_path):
    report = XtbReportReader().read(report_path)
    earlier = report.model_copy(
        update={
            "as_of": datetime(2026, 9, 15, tzinfo=UTC),
            "fingerprint": "f" * 64,
        }
    )
    repository.import_report(earlier)
    # Net quantity is unchanged at Oct 1, but a Sep 10 purchase changes the
    # already stored Sep 15 snapshot until its Sep 20 sale. Check all dates.
    buy = report.cash_events[1].model_copy(
        update={
            "operation_id": "late-discovered-buy",
            "quantity": Decimal("0.25"),
            "amount": Decimal("-25"),
            "comment": "OPEN BUY 0.25 @ 100",
            "occurred_at": datetime(2026, 9, 10, tzinfo=UTC),
        }
    )
    sale = report.cash_events[3].model_copy(
        update={
            "operation_id": "late-discovered-sale",
            "occurred_at": datetime(2026, 9, 20, tzinfo=UTC),
        }
    )
    report.cash_events.extend([buy, sale])
    with pytest.raises(ImportValidationError, match="invalidate a stored portfolio snapshot"):
        repository.import_report(report)
    assert count(repository, ImportBatch) == 1
    assert count(repository, Transaction) == 4
